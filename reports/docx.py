from docx import Document
import json

def generate_docx(meeting, output_path: str):
    doc = Document()
    
    doc.add_heading("MeetIntel AI - Minutes of Meeting", 0)
    
    doc.add_heading("Meeting Details", level=1)
    doc.add_paragraph(f"Title: {meeting.title}")
    doc.add_paragraph(f"Date: {meeting.date.strftime('%Y-%m-%d %H:%M')}")
    duration_str = f"{int(meeting.duration)//60}m {int(meeting.duration)%60}s" if meeting.duration else "N/A"
    doc.add_paragraph(f"Duration: {duration_str}")
    
    doc.add_heading("Executive Summary", level=1)
    doc.add_paragraph(meeting.executive_summary if meeting.executive_summary else "N/A")
    
    doc.add_heading("Participants", level=1)
    if meeting.participants:
        for p in meeting.participants:
            doc.add_paragraph(f"{p.name} ({p.speaker_id})", style='List Bullet')
    else:
        doc.add_paragraph("None")
    
    doc.add_heading("Action Items", level=1)
    if meeting.action_items:
        for ai in meeting.action_items:
            p = doc.add_paragraph(style='List Bullet')
            p.add_run(f"{ai.task} ").bold = True
            p.add_run(f"(Assigned to: {ai.assigned_to}, Deadline: {ai.deadline}, Priority: {ai.priority})")
    else:
        doc.add_paragraph("None")
        
    doc.add_heading("Decisions", level=1)
    if meeting.decisions:
        for d in meeting.decisions:
            p = doc.add_paragraph(style='List Bullet')
            p.add_run(f"{d.decision}").bold = True
    else:
        doc.add_paragraph("None")
        
    doc.save(output_path)
    return output_path
